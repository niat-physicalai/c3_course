#!/usr/bin/env python3
"""Step 0 (one-time) — inspect the template .pptx and draft deck/template.map.yaml.

  python tools/extract_template.py            # writes the draft map (won't overwrite) + inventory
  python tools/extract_template.py --force    # overwrite an existing map
  python tools/extract_template.py --sample   # also build deck/build/_layout-sample.pptx (one slide per kind)

The map tells build_deck.py which layout (or which reference slide to clone) to use
for each slide kind, and where title / body / visual go.
"""
from __future__ import annotations

import _venv  # noqa: F401  (re-runs under .venv if needed)

import sys

import yaml
from pptx import Presentation
from pptx.util import Emu

from common import ROOT, load_config

KINDS = {   # kind → layout-name keywords, in preference order
    "title":      [["title slide"], ["title"]],
    "agenda":     [["title and content"], ["title and body"], ["content"]],
    "section":    [["section header"], ["section"], ["title only"]],
    "image-text": [["two content"], ["comparison"], ["title and two"], ["picture with caption"], ["title only"]],
    "image-full": [["title only"], ["picture"], ["blank"]],
    "two-images": [["title only"], ["comparison"], ["blank"]],
    "bullets":    [["title and content"], ["title and body"], ["content"]],
    "key-idea":   [["title only"], ["section"], ["big number"], ["blank"]],
    "recap":      [["title and content"], ["title and body"], ["content"]],
}
IN = 914400


def inches(v):
    return round(Emu(v).inches, 2)


def find_layout(prs, prefs):
    layouts = list(prs.slide_layouts)
    for words in prefs:
        for lay in layouts:
            n = lay.name.lower()
            if all(w in n for w in words):
                return lay
    return layouts[min(5, len(layouts) - 1)]


def ph_by(lay, *types):
    out = []
    for ph in lay.placeholders:
        t = str(ph.placeholder_format.type).split(".")[-1].split(" ")[0]
        if any(t.startswith(x) for x in types):
            out.append(ph)
    return out


def draft_map(prs):
    W, H = inches(prs.slide_width), inches(prs.slide_height)
    m, top = 0.5, 1.45
    full_box = [m, top, round(W - 2 * m, 2), round(H - top - 0.55, 2)]
    half = round((W - 3 * m) / 2, 2)
    kinds = {}
    for kind, prefs in KINDS.items():
        lay = find_layout(prs, prefs)
        e = {"layout": lay.name}
        t = ph_by(lay, "TITLE", "CENTER_TITLE")
        if t:
            e["title"] = t[0].placeholder_format.idx
        bodies = [p for p in ph_by(lay, "BODY", "OBJECT", "SUBTITLE", "PICTURE")
                  if p.placeholder_format.idx not in (e.get("title"),)]
        bodies.sort(key=lambda p: (p.left or 0))
        if kind in ("title", "section"):
            if bodies:
                e["subtitle"] = bodies[0].placeholder_format.idx
        elif kind in ("agenda", "bullets", "recap"):
            if bodies:
                e["body"] = bodies[0].placeholder_format.idx
            else:
                e["body_box"] = full_box
        elif kind == "image-text":
            if len(bodies) >= 2:
                e["body"], e["visual"] = bodies[0].placeholder_format.idx, bodies[1].placeholder_format.idx
            else:
                e["body_box"] = [m, top, half, full_box[3]]
                e["visual_box"] = [round(2 * m + half, 2), top, half, full_box[3]]
        elif kind == "image-full":
            e["visual_box"] = [m, top, full_box[2], round(full_box[3] - 0.6, 2)]
            e["caption_box"] = [m, round(top + full_box[3] - 0.5, 2), full_box[2], 0.45]
        elif kind == "two-images":
            e["visual_box"] = [m, top, half, round(full_box[3] - 0.6, 2)]
            e["visual2_box"] = [round(2 * m + half, 2), top, half, round(full_box[3] - 0.6, 2)]
            e["caption_box"] = [m, round(top + full_box[3] - 0.5, 2), full_box[2], 0.45]
        elif kind == "key-idea":
            e["body_box"] = [1.0, round(H * 0.3, 2), round(W - 2.0, 2), round(H * 0.4, 2)]
        kinds[kind] = e
    return {
        "slide_size_in": [W, H],
        "keep_template_slides": False,
        "fonts": {"title_pt": 32, "body_pt": 20, "key_idea_pt": 32, "caption_pt": 12,
                  "body_font": None, "caption_color": "666666"},
        "kinds": kinds,
    }


def inventory(prs):
    out = ["# Template inventory", "", f"Slide size: {inches(prs.slide_width)} × {inches(prs.slide_height)} in", "",
           "## Layouts (placeholder idx · type · name · box in inches)"]
    for i, lay in enumerate(l for mst in prs.slide_masters for l in mst.slide_layouts):
        out.append(f"### [{i}] {lay.name}")
        for ph in lay.placeholders:
            pf = ph.placeholder_format
            out.append(f"- idx {pf.idx} · {str(pf.type).split('.')[-1]} · {ph.name} · "
                       f"[{inches(ph.left or 0)}, {inches(ph.top or 0)}, {inches(ph.width or 0)}, {inches(ph.height or 0)}]")
    out += ["", "## Existing slides (for `prototype:` mode — clone a designed slide)"]
    for n, s in enumerate(prs.slides, 1):
        out.append(f"### slide {n} — layout '{s.slide_layout.name}'")
        for sh in s.shapes:
            txt = (sh.text_frame.text[:40].replace("\n", " ") if sh.has_text_frame else "")
            out.append(f"- \"{sh.name}\" {sh.shape_type} [{inches(sh.left or 0)}, {inches(sh.top or 0)}, "
                       f"{inches(sh.width or 0)}, {inches(sh.height or 0)}] {txt}")
    return "\n".join(out)


def main():
    cfg = load_config()
    tpath = ROOT / cfg["template"]
    if not tpath.exists():
        sys.exit(f"Template not found: {tpath}")
    prs = Presentation(tpath)
    mpath = ROOT / cfg.get("template_map", "deck/template.map.yaml")
    mpath.parent.mkdir(parents=True, exist_ok=True)
    (mpath.parent / "template-inventory.md").write_text(inventory(prs), encoding="utf-8")
    print(f"✓ inventory → {(mpath.parent / 'template-inventory.md').relative_to(ROOT)}")
    if mpath.exists() and "--force" not in sys.argv:
        print(f"• {mpath.relative_to(ROOT)} exists — kept (use --force to regenerate)")
    else:
        header = ("# Slide-kind → template mapping. Edit freely; build_deck.py reads this.\n"
                  "# Per kind use EITHER  layout: <name>  (+ placeholder idx for title/subtitle/body/visual)\n"
                  "#   OR  prototype: <slide number in template>  (+ shape NAMES for title/body/visual) to clone a designed slide.\n"
                  "# *_box values are [left, top, width, height] in inches and override placeholders.\n")
        mpath.write_text(header + yaml.safe_dump(draft_map(prs), sort_keys=False, allow_unicode=True), encoding="utf-8")
        print(f"✓ draft map → {mpath.relative_to(ROOT)}  (review it!)")
    if "--sample" in sys.argv:
        from build_deck import build
        spec = {"module": "_sample", "title": "Layout sample", "slides": [
            {"id": f"s{i:02d}", "kind": k, "title": f"Kind: {k}", "subtitle": "subtitle",
             "bullets": ["First short point", "Second short point", "Third point"],
             "key_idea": "One sentence that matters.",
             "visual": {"id": "SAMPLE", "status": "todo", "caption": "caption"},
             "visual2": {"id": "SAMPLE2", "status": "todo"}} for i, k in enumerate(KINDS, 1)]}
        out = ROOT / cfg.get("output_dir", "deck/build") / "_layout-sample.pptx"
        build(cfg, spec, out, ROOT)
        print(f"✓ sample → {out.relative_to(ROOT)}  (render it with python3 tools/render_preview.py)")


if __name__ == "__main__":
    main()
