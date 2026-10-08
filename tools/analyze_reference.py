#!/usr/bin/env python3
"""Step 0 (one-time) — learn the house style from the Applied IoT reference pairs.

  python tools/analyze_reference.py

For every reference pair (markdown folder ↔ pptx) it measures:
  • slide count, words/slide, visuals/slide, layouts used, fonts/sizes used
  • compression: reading-material words → slide words
  • which markdown section each slide was distilled from (word-overlap match)
Writes deck/style-profile.json (+ .md for humans / Claude to read).
"""
from __future__ import annotations

import _venv  # noqa: F401  (re-runs under .venv if needed)

import re
import statistics as st
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from common import ROOT, load_config, write_json
from digest_module import parse_unit

WORD = re.compile(r"[a-z][a-z0-9+-]{2,}")


def toks(t: str) -> set[str]:
    return set(WORD.findall(t.lower()))


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def slide_stats(slide):
    title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
    body, pics, gifs, fonts, sizes, tables, charts = [], 0, 0, Counter(), Counter(), 0, 0
    for sh in iter_shapes(slide.shapes):
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(sh, "image", None) is not None and sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER:
            try:
                pics += 1
                if sh.image.ext.lower() == "gif":
                    gifs += 1
            except Exception:
                pass
        if getattr(sh, "has_table", False) and sh.has_table:
            tables += 1
        if getattr(sh, "has_chart", False) and sh.has_chart:
            charts += 1
        if sh.has_text_frame and sh != slide.shapes.title:
            body.append(sh.text_frame.text)
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        fonts[r.font.name] += 1
                    if r.font.size:
                        sizes[round(r.font.size.pt)] += 1
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    text = "\n".join(body)
    return {"layout": slide.slide_layout.name, "title": title, "body_words": len(text.split()),
            "body": text[:400], "pictures": pics, "gifs": gifs, "tables": tables, "charts": charts,
            "notes_words": len(notes.split()), "_fonts": fonts, "_sizes": sizes}


def analyze_pair(pair):
    prs = Presentation(ROOT / pair["pptx"])
    if pair.get("markdown_files"):          # explicit files / globs (one deck per reading-material unit)
        paths = []
        for pat in pair["markdown_files"]:
            paths += sorted(ROOT.glob(pat))
    else:
        paths = sorted((ROOT / pair["markdown_dir"]).glob("*.md"))
    if not paths:
        raise SystemExit(f"No markdown found for reference pair '{pair['name']}'")
    units = [parse_unit(p) for p in paths]
    sections = [(u["id"], s["heading"], toks(s["heading"] + " " + s["text"])) for u in units for s in u["sections"]]
    md_words = sum(u["words"] for u in units)

    slides, fonts, sizes = [], Counter(), Counter()
    for i, s in enumerate(prs.slides, 1):
        d = slide_stats(s)
        fonts.update(d.pop("_fonts"))
        sizes.update(d.pop("_sizes"))
        st_toks = toks(d["title"] + " " + d["body"])
        best = max(sections, key=lambda x: len(st_toks & x[2]) / (len(st_toks | x[2]) or 1), default=None)
        if best and st_toks:
            score = len(st_toks & best[2]) / (len(st_toks) or 1)
            d["source_guess"] = f"{best[0]} › {best[1]} ({score:.0%} of slide words found)"
        d["n"] = i
        slides.append(d)

    content = [s for s in slides if s["body_words"] or s["pictures"]]
    slide_words = sum(s["body_words"] + len(s["title"].split()) for s in slides)
    return {
        "name": pair["name"], "pptx": pair["pptx"], "slide_count": len(slides),
        "slide_size_in": [round(prs.slide_width / 914400, 2), round(prs.slide_height / 914400, 2)],
        "md_words": md_words, "slide_words": slide_words,
        "compression_ratio": round(md_words / slide_words, 1) if slide_words else None,
        "md_words_per_slide": round(md_words / len(slides)) if slides else None,
        "avg_body_words": round(st.mean([s["body_words"] for s in content]), 1) if content else 0,
        "visual_ratio": round(sum(1 for s in content if s["pictures"]) / len(content), 2) if content else 0,
        "gif_count": sum(s["gifs"] for s in slides),
        "layouts_used": Counter(s["layout"] for s in slides).most_common(),
        "fonts": fonts.most_common(6), "font_sizes_pt": sorted(sizes.most_common(8)),
        "slides": slides,
    }


def main():
    cfg = load_config()
    pairs = []
    for pair in cfg.get("reference_pairs") or []:
        if not (ROOT / pair["pptx"]).exists():
            print(f"! skip {pair['name']}: {pair['pptx']} not found")
            continue
        pairs.append(analyze_pair(pair))
        print(f"✓ {pair['name']}: {pairs[-1]['slide_count']} slides, compression {pairs[-1]['compression_ratio']}x")
    if not pairs:
        raise SystemExit("No reference pptx found — put the Applied IoT decks where deck.config.yaml points.")

    agg = {
        "avg_slide_count": round(st.mean(p["slide_count"] for p in pairs), 1),
        "avg_compression_ratio": round(st.mean(p["compression_ratio"] for p in pairs if p["compression_ratio"]), 1),
        "avg_md_words_per_slide": round(st.mean(p["md_words_per_slide"] for p in pairs if p["md_words_per_slide"])),
        "avg_body_words_per_slide": round(st.mean(p["avg_body_words"] for p in pairs), 1),
        "avg_visual_ratio": round(st.mean(p["visual_ratio"] for p in pairs), 2),
    }
    out = Path(ROOT / cfg.get("style_profile", "deck/style-profile.json"))
    write_json(out, {"aggregate": agg, "pairs": pairs})

    md = ["# Style profile (learned from reference decks)", "",
          "Use these as targets, not hard rules. Content still comes ONLY from the module being built.", "",
          "| metric | value |", "|---|---|"] + [f"| {k} | {v} |" for k, v in agg.items()] + [""]
    for p in pairs:
        md += [f"## {p['name']} — {p['slide_count']} slides, {p['slide_size_in'][0]}×{p['slide_size_in'][1]} in",
               f"Reading {p['md_words']} words → slides {p['slide_words']} words (×{p['compression_ratio']}). "
               f"Visual ratio {p['visual_ratio']}, GIFs {p['gif_count']}.",
               f"Layouts: {p['layouts_used']}", f"Fonts: {p['fonts']}  Sizes(pt,count): {p['font_sizes_pt']}", "",
               "| # | layout | title | body w | pics | gifs | distilled from |", "|---|---|---|---|---|---|---|"]
        for s in p["slides"]:
            md.append(f"| {s['n']} | {s['layout']} | {s['title'][:50]} | {s['body_words']} | {s['pictures']} | "
                      f"{s['gifs']} | {s.get('source_guess', '')} |")
        md.append("")
    out.with_suffix(".md").write_text("\n".join(md), encoding="utf-8")
    print(f"✓ wrote {out.relative_to(ROOT)} and .md")


if __name__ == "__main__":
    main()
